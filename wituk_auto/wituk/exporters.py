"""오픈마켓 대량등록 엑셀 만들기.

방식: 마켓에서 내려받은 '진짜 양식 파일'을 열어, 머리글 이름이 설정(config/markets/*.json)의
"columns" 와 같은 칸에 값을 채워 새 파일로 저장합니다. 양식 원본은 건드리지 않습니다.
양식 파일이 없으면 설정에 적힌 열 이름으로 기본 엑셀을 만듭니다.

columns 값 문법:
  "clean_title"      → 상품 정보에서 가져옴(아래 FIELDS 참고)
  "=신상품"          → 항상 같은 글자
  "setting:as_phone" → settings.json 의 값
"""
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .pricing import landed_cost, profit


def _norm(s):
    return str(s or "").replace(" ", "").replace("\n", "").replace("*", "").replace("(필수)", "").lower()


def field_values(p, market_cfg, cfg):
    sep = market_cfg.get("multi_sep", "|")
    if market_cfg.get("image_mode", "url") == "file":
        imgs = [Path(x).name for x in p.local_images]
    else:
        imgs = [x for x in p.images if str(x).startswith("http")] or [Path(x).name for x in p.local_images]
    opt_sep = market_cfg.get("option_sep", ",")
    return {
        "clean_title": p.clean_title,
        "title": p.title,
        "sale_price": p.sale_price,
        "stock": min(p.stock or cfg.get("default_stock", 99), cfg.get("max_stock", 999)),
        "category": p.category,
        "category_id": cfg.get("category_map", {}).get(p.category, ""),
        "brand": p.brand,
        "image_main": imgs[0] if imgs else "",
        "images_extra": sep.join(imgs[1:10]),
        "detail_html": p.detail_html or "".join(f'<img src="{u}"><br>' for u in imgs),
        "keywords": ",".join(p.keywords),
        "seller_code": f"{p.source[:2].upper()}-{p.source_id}",
        "option_group": market_cfg.get("option_group_name", "옵션") if p.options else "",
        "option_names": opt_sep.join(o.name for o in p.options),
        "option_prices": opt_sep.join(str(o.extra_price) for o in p.options),
        "option_stocks": opt_sep.join(str(o.stock or cfg.get("default_stock", 99)) for o in p.options),
        "source_url": p.url,
        "cost": p.cost,
    }


def _value(expr, fields, cfg):
    if not isinstance(expr, str):
        return expr
    if expr.startswith("="):
        return expr[1:]
    if expr.startswith("setting:"):
        return cfg.get(expr[8:], "")
    return fields.get(expr, "")


def export_market(products, market_cfg, cfg, out_path, template=None):
    cols = market_cfg["columns"]
    if template and Path(template).exists():
        wb = load_workbook(template)
        ws = wb[market_cfg["sheet"]] if market_cfg.get("sheet") in wb.sheetnames else wb.active
        hr = market_cfg.get("header_row", 1)
        start = market_cfg.get("data_start_row", hr + 1)
        header = {_norm(c.value): c.column for c in ws[hr] if c.value}
        col_idx, missing = {}, []
        for name, expr in cols.items():
            idx = header.get(_norm(name))
            if idx:
                col_idx[idx] = expr
            else:
                missing.append(name)
        # 양식에 이미 들어 있던 예시 데이터는 지운다
        for r in range(start, ws.max_row + 1):
            for c in col_idx:
                ws.cell(r, c).value = None
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = market_cfg.get("sheet") or "등록"
        for i, name in enumerate(cols, 1):
            ws.cell(1, i, name).font = Font(bold=True)
        col_idx = {i: expr for i, expr in enumerate(cols.values(), 1)}
        start, missing = 2, []

    r = start
    for p in products:
        f = field_values(p, market_cfg, cfg)
        for c, expr in col_idx.items():
            ws.cell(r, c).value = _value(expr, f, cfg)
        r += 1
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return missing


def export_report(products, cfg, market_names, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "결과"
    head = ["상태", "소싱처", "상품번호", "원본 상품명", "정리된 상품명", "도매가", "원가합계",
            "판매가", "예상이익(1개)", "재고", "옵션수", "이미지수", "경고", "링크"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    red = PatternFill("solid", fgColor="F8CBAD")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    mk = market_names[0] if market_names else None
    for p in products:
        ws.append([
            "제외" if p.blocked else ("확인필요" if p.flags else "등록"),
            p.source, p.source_id, p.title, p.clean_title, p.cost, int(landed_cost(p, cfg)),
            p.sale_price, profit(p.sale_price, p, cfg, mk) if p.sale_price else "",
            p.stock, len(p.options), len(p.local_images) or len(p.images), ", ".join(p.flags), p.url,
        ])
        if p.blocked or p.flags:
            for c in ws[ws.max_row]:
                c.fill = red if p.blocked else yellow
    for col, w in zip("ABCDEFGHIJKLMN", [8, 10, 12, 40, 40, 9, 9, 9, 11, 7, 7, 8, 30, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    info = wb.create_sheet("설정값")
    info.append(["만든 시각", datetime.now().strftime("%Y-%m-%d %H:%M")])
    for k in ("margin_rate", "default_fee_rate", "my_shipping_fee", "free_shipping", "round_unit"):
        info.append([k, json.dumps(cfg.get(k), ensure_ascii=False)])
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
