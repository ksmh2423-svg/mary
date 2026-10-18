"""도매처에서 내려받은 엑셀/CSV 상품목록 불러오기.

도매처마다 열 이름이 달라서, 아래 별칭 중 하나와 일치하는 열을 찾아 씁니다.
맞는 열이 없으면 config/settings.json 의 "excel_import_columns" 에 추가하면 됩니다.
"""
import csv
from pathlib import Path

from openpyxl import load_workbook

from ..models import Option, Product

DEFAULT_ALIASES = {
    "source_id": ["상품번호", "상품코드", "품번", "자체상품코드", "no", "id", "item_no"],
    "title": ["상품명", "제품명", "품명", "title", "name"],
    "cost": ["도매가", "공급가", "공급가격", "판매가", "가격", "price", "cost"],
    "stock": ["재고", "재고수량", "수량", "stock", "qty"],
    "category": ["카테고리", "분류", "category"],
    "brand": ["브랜드", "제조사", "brand"],
    "images": ["대표이미지", "이미지", "이미지url", "썸네일", "image", "thumb"],
    "extra_images": ["추가이미지", "추가이미지url", "서브이미지"],
    "detail_html": ["상세설명", "상세페이지", "상세html", "detail"],
    "options": ["옵션", "옵션정보", "option"],
    "source_shipping": ["배송비", "택배비", "shipping"],
    "url": ["상품url", "상품링크", "url", "link"],
    "currency": ["통화", "currency"],
}


def _norm(s):
    return str(s or "").replace(" ", "").replace("_", "").lower()


def _to_int(v):
    if v is None or v == "":
        return 0
    try:
        return int(float(str(v).replace(",", "").replace("원", "").strip()))
    except ValueError:
        return 0


def _read_rows(path):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "cp949"):
            try:
                with open(path, newline="", encoding=enc) as f:
                    return [r for r in csv.reader(f)]
            except UnicodeDecodeError:
                continue
        raise ValueError(f"CSV 인코딩을 알 수 없습니다: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _find_header(rows, aliases):
    """앞쪽 10줄 중 상품명 열이 있는 줄을 머리글로 본다."""
    title_keys = {_norm(a) for a in aliases["title"]}
    for i, row in enumerate(rows[:10]):
        if any(_norm(c) in title_keys for c in row):
            return i
    raise ValueError("상품명 열을 찾지 못했습니다. 머리글 이름을 확인해 주세요.")


def parse_options(text):
    """'블랙/L(+1000)|화이트/M' 또는 줄바꿈 구분 옵션 문자열을 Option 목록으로."""
    opts = []
    if not text:
        return opts
    for part in str(text).replace("\n", "|").replace(",", "|").split("|"):
        part = part.strip()
        if not part:
            continue
        extra = 0
        if "(+" in part and part.endswith(")"):
            part, num = part.rsplit("(+", 1)
            extra = _to_int(num.rstrip(")"))
        elif "(-" in part and part.endswith(")"):
            part, num = part.rsplit("(-", 1)
            extra = -_to_int(num.rstrip(")"))
        opts.append(Option(name=part.strip(), extra_price=extra, stock=0))
    return opts


def load_products(path, extra_aliases=None, source_name="excel"):
    aliases = {k: list(v) for k, v in DEFAULT_ALIASES.items()}
    for k, v in (extra_aliases or {}).items():
        aliases.setdefault(k, []).extend(v)

    rows = _read_rows(path)
    h = _find_header(rows, aliases)
    header = [_norm(c) for c in rows[h]]
    col = {}
    for key, names in aliases.items():
        for n in names:
            if _norm(n) in header:
                col[key] = header.index(_norm(n))
                break

    def get(row, key):
        i = col.get(key)
        return row[i] if i is not None and i < len(row) else None

    products = []
    for n, row in enumerate(rows[h + 1:], start=1):
        title = get(row, "title")
        if not title or not str(title).strip():
            continue
        imgs = [s.strip() for s in str(get(row, "images") or "").replace("\n", "|").split("|") if s.strip()]
        imgs += [s.strip() for s in str(get(row, "extra_images") or "").replace("\n", "|").split("|") if s.strip()]
        base = Path(path).parent  # 상대 경로 이미지는 엑셀 파일 옆 기준
        imgs = [i if i.startswith("http") or Path(i).is_absolute() else str(base / i) for i in imgs]
        products.append(Product(
            source=source_name,
            source_id=str(get(row, "source_id") or f"row{n}"),
            title=str(title).strip(),
            cost=_to_int(get(row, "cost")),
            stock=_to_int(get(row, "stock")),
            category=str(get(row, "category") or ""),
            brand=str(get(row, "brand") or ""),
            images=imgs,
            detail_html=str(get(row, "detail_html") or ""),
            options=parse_options(get(row, "options")),
            source_shipping=_to_int(get(row, "source_shipping")),
            url=str(get(row, "url") or ""),
            currency=str(get(row, "currency") or "KRW").upper() or "KRW",
        ))
    return products
